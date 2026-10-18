"""
Ubah Kamus_Slang_Hook_Konten.xlsx menjadi src/data/kamus-hook.js.

Dipakai services/kamus-hook.js untuk memberi AI penyusun brief bahan slang
dan rumus hook. Hasilnya modul JS (bukan JSON yang dibaca saat jalan) supaya
ikut terbundel di Vercel.

Perbarui kamus: ganti berkas xlsx di folder ini, lalu
    python data/bangun-kamus.py
dari folder lokalboost-api. Sheet "Data Training" sengaja tidak dibaca: isinya
rumus turunan dari "Kamus Slang".
"""
import json
import pathlib
import sys

import openpyxl

akar = pathlib.Path(__file__).resolve().parent
sumber = akar / 'Kamus_Slang_Hook_Konten.xlsx'
tujuan = akar.parent / 'src' / 'data' / 'kamus-hook.js'

wb = openpyxl.load_workbook(sumber, data_only=False)


def pisah(teks):
    return [t.strip() for t in str(teks or '').split(',') if t.strip()]


slang = []
for baris in wb['Kamus Slang'].iter_rows(min_row=2, values_only=True):
    if not baris[0]:
        continue
    (id_, kata, variasi, kategori, arti, fungsi, contoh, niche, tone, tren, aman, catatan) = baris[:12]
    slang.append(
        {
            'id': id_,
            'kata': kata,
            'variasi': pisah(variasi),
            'kategori': kategori,
            'arti': arti,
            'fungsi': fungsi,
            'contoh': contoh,
            'niche': pisah(niche),
            'tone': pisah(tone),
            'tren': tren,
            'keamanan': aman,
            'catatan': catatan or None,
        }
    )

template = []
for baris in wb['Template Hook'].iter_rows(min_row=2, values_only=True):
    if not baris[0]:
        continue
    tipe, rumus, isian, contoh = baris[:4]
    template.append({'tipe': tipe, 'rumus': rumus, 'isian': isian, 'contoh': contoh})

if len(slang) < 50 or len(template) < 5:
    sys.exit(f'Isi kamus mencurigakan: {len(slang)} slang, {len(template)} template. Dibatalkan.')

tujuan.parent.mkdir(parents=True, exist_ok=True)
isi = (
    '// DIBUAT OTOMATIS oleh data/bangun-kamus.py dari Kamus_Slang_Hook_Konten.xlsx. Jangan disunting tangan.\n'
    f'export const SLANG = {json.dumps(slang, ensure_ascii=False, indent=1)}\n\n'
    f'export const TEMPLATE_HOOK = {json.dumps(template, ensure_ascii=False, indent=1)}\n'
)
tujuan.write_text(isi, encoding='utf-8')
print(f'{len(slang)} slang, {len(template)} template -> {tujuan}')
